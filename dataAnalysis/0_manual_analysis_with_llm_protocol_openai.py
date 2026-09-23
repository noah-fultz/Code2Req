"""
Uses separate output/checkpoint locations so this script doesn't clobber the
backed script's results:
  - Output workbook: Manual_analysis_llm_protocol_openai.xlsx
  - Checkpoints:      llm_checkpoints_protocol_openai/{dataset}_{variant}.csv

Auth: set the OPENAI_API_KEY environment variable (or edit OPENAI_API_KEY
below) before running. Get a key from https://platform.openai.com/api-keys.

Running multiple times (repeat measurement): pass --runs N (e.g. --runs 3) to
run the whole LLM pipeline N times per dataset/variant. All runs are appended
to the same checkpoint CSV / output workbook, tagged with a "run" column.
--agreement then computes Cohen's kappa and % agreement separately for each
run found in the checkpoints, plus the mean and sample standard deviation
across runs. The standard deviation describes run-to-run LLM variability.

Run analysis:
python 0_manual_analysis_with_llm_protocol_openai.py --runs 3 --agreement
"""

import argparse
import csv
import os
import time
from pathlib import Path
from typing import Dict, List, Literal, Optional, Tuple

import pandas as pd
from openai import OpenAI
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.dataframe import dataframe_to_rows
from pydantic import BaseModel, Field, create_model
from sklearn.metrics import cohen_kappa_score

# ============================================================
# Configuration
# ============================================================
SCRIPT_DIR = Path(__file__).resolve().parent
EXCEL_PATH = SCRIPT_DIR / "Manual_analysis.xlsx"
OUTPUT_XLSX = SCRIPT_DIR / "Manual_analysis_llm_protocol_openai.xlsx"
CHECKPOINT_DIR = SCRIPT_DIR / "llm_checkpoints_protocol_openai"

DATASETS = ["eANCI", "SMOS", "eTour", "iTrust"]
VARIANTS = [
    "Code_RAG_GPT",
    "Code_GPT",
    "Feature_RAG_GPT",
    "Feature_GPT",
    "Code_RAG_Qwen",
    "Code_Qwen",
    "Feature_RAG_Qwen",
    "Feature_Qwen",
]

# Reads the key from the OPENAI_API_KEY env var by default; fall back to a
# literal string here only if you don't want to set the env var.
OPENAI_API_KEY = os.environ.get("OPENAI_API_KEY", "")
OPENAI_MODEL = "gpt-5.6-sol"

RETRIES = 3
RETRY_BACKOFF_SECONDS = 15


# ============================================================
# LLM client (OpenAI API-key auth)
# ============================================================
def get_client() -> OpenAI:
    if not OPENAI_API_KEY:
        raise RuntimeError(
            "No OpenAI API key found. Set the OPENAI_API_KEY environment variable "
            "(or edit OPENAI_API_KEY in this script)."
        )
    return OpenAI(api_key=OPENAI_API_KEY)


def get_model_name() -> str:
    return OPENAI_MODEL


# ============================================================
# Response schema (mirrors the columns filled in during manual analysis).
# Descriptions here are intentionally minimal - the full decision rules live
# in ANALYSIS_INSTRUCTIONS (the verbatim protocol text) below, not here.
# Only the 7 fields assessable without source-code access are included;
# "Key concept code class" and "Recommended addition" require inspecting the
# project's source code, which isn't provided to the LLM in this pipeline.
# ============================================================
class RequirementAnalysis(BaseModel):

    syntactic_adherence: Literal["Yes", "No"] = Field(
        description=(
            "Syntactic adherence, per protocol criterion 1. "
            "Judge whether the generated requirement follows the characteristic structural "
            "and presentation pattern of the project's gold requirements. Partial use of the "
            "same structure is not sufficient when multiple characteristic elements are missing. "
            "Judge structure only, not semantic quality or completeness."
        )
    )

    semantic_adherence: Literal["Yes", "No"] = Field(
        description=(
            "Semantic adherence, per protocol criterion 2. "
            "Judge whether the requirement's main intent is a domain-level goal, capability, "
            "rule, constraint, or behavior of the kind represented in the gold set. "
            "Actor names, domain entities, or use-case formatting do not by themselves establish "
            "adherence. Internal implementation, persistence, framework, configuration, or helper "
            "operations are not domain-level merely because they are written as requirements. "
            "A lookup or validation is domain-level (Yes) whenever its outcome is consumed directly by "
            "an actor - returning a value, granting or denying access, or redirecting on failure all "
            "count as actor-facing, even when phrased in technical or procedural terms. Judge No only "
            "when the operation purely supports other internal code with no actor-facing outcome of "
            "its own."
        )
    )

    key_concept_manually_matched_to: List[str] = Field(
        default_factory=list,
        description=(
            "Key concept manually matched to, per protocol criterion 4. "
            "Match gold requirements that express the same core domain functionality or user goal, based on "
            "the main action and its primary object or intended outcome - not shared terminology, data, "
            "topic, prerequisites, or supporting/validation steps alone. A validation or error-handling step "
            "(e.g. 'the given id is invalid') does not match every gold exception condition that validates "
            "the same kind of id in an unrelated use case, and a narrow single-category view does not match "
            "a broad multi-category gold requirement merely because that category is one of several items it "
            "lists - only when retrieving that one item is itself the broad requirement's distinct core goal. "
            "Differences in actor, abstraction, or detail do not prevent a match. A narrower requirement may "
            "match a broader gold requirement when it directly expresses a substantive part of the same "
            "functional goal; incidental enabling steps do not. A broad requirement may match multiple gold "
            "requirements only when it clearly includes each of their core functionalities. Return unique "
            "matching IDs, or [] if unmatched."
        ),
    )

    secondary_concept_match: List[Literal["Yes", "Partly", "No"]] = Field(
        default_factory=list,
        description=(
            "Secondary concept match, per protocol criterion 5. "
            "For each mapped gold requirement, compare only its specified actors/roles and additional "
            "domain actions beyond the shared key concept - not incidental steps needed only to reach or "
            "perform that action (e.g. browsing or selecting from a list before acting), which are not "
            "separate secondary concepts. Do not infer omitted concepts. "
            "Yes means all are covered; Partly requires at least one clearly covered and at least one "
            "clearly missing or different; No means none are covered. Partly is not an uncertainty label. "
            "Equivalent wording counts, but a generic actor does not automatically equal a specific role - "
            "judge by role/responsibility rather than exact wording, though a functionally narrower or "
            "different actor is still a difference. If the gold requirement specifies alternative or "
            "optional secondary actions (e.g. 'via X or Y'), covering only one alternative is partial, "
            "not full, coverage. "
            "Technical mechanisms, prescribed implementation details, and specific enumerated values "
            "belong under Details match, not here. "
            "One label per matched ID, in the same order; [] if unmatched."
        ),
    )

    details_match: List[Literal["Yes", "Partly", "No"]] = Field(
        default_factory=list,
        description=(
            "Details match, per protocol criterion 6. "
            "Compare the mapped gold requirement's specified fields, values, formats, constraints, "
            "conditions, validations, and prescribed steps. Do not infer unstated details. When gold "
            "names a specific field, code, format (e.g. DD-MM-YY), or enumerated item, it must be named or clearly "
            "equivalent in the generated requirement; a generic restatement (e.g. 'validates inputs') "
            "does not preserve it and counts as missing, not partial. Yes requires all specified details "
            "to be preserved. Partly is limited to minor omissions or differences while all "
            "defining constraints and required behavior remain intact. Use No when defining details "
            "are missing, generalized, or changed. "
            "Additional generated details do not compensate for missing gold details. "
            "One label per matched ID, in the same order; [] if unmatched."
        ),
    )

    abstraction_level: Literal["Equal", "Lower", "Higher"] = Field(
        description=(
            "Abstraction level, per protocol criterion 7. "
            "Judge the breadth of the main domain functionality. For mapped requirements, compare "
            "primarily with the mapped gold requirement(s); otherwise compare with the project's typical "
            "gold-requirement scope. Lower is the default for anything short of a complete, self-contained "
            "goal: a stage, subfunction, exception, or supporting operation of a gold-level goal; covering "
            "only one of several actor types, cases, or branches that a gold requirement bundles together; "
            "or a bare single data operation (create, list, search, retrieve, update, or delete on one "
            "entity or field) with no further distinguishing actor goal. Equal applies only when the "
            "requirement's action itself matches the completeness and standalone status that gold "
            "requirements of this kind typically have - not merely because the requirement reads as "
            "complete in isolation. Higher means multiple gold-level goals are combined."
        )
    )

    detail_level: Literal["Equal", "Lower", "Higher"] = Field(
        description=(
            "Detail level, per protocol criterion 8. "
            "Judge how much of the comparison gold requirement's (or, for unmatched requirements, the "
            "typical gold requirement's) concrete domain-specific data - exact enumerated values, named "
            "fields, data-format references, thresholds, algorithms - is actually reproduced."
            "Lower: applies whenever only some, or none, of those concrete data points are reproduced, "
            "even when the requirement is long or elaborated with generic validation, logging, or "
            "error-handling steps, since that procedural elaboration is implementation boilerplate, not "
            "concrete domain detail, and does not raise the rating. Reserve Equal for requirements that "
            "name nearly all of the same concrete domain-specific data points as the comparison gold "
            "requirement, not merely one or two of several. Higher means materially more concrete "
            "domain-specific information than that. A requirement may have Details match = No while "
            "still having Equal or Higher detail level."
        )
    )

ANALYSIS_FIELDS = list(RequirementAnalysis.model_fields.keys())

# Batch schema: same fields as RequirementAnalysis, plus a requirement_id so the LLM
# can return one entry per requirement in a single response (used by --batch mode).
RequirementAnalysisItem = create_model(
    "RequirementAnalysisItem",
    requirement_id=(str, Field(description="The id of the requirement this analysis corresponds to, exactly as given in the input (e.g. 'code_req_42').")),
    **{name: (field.annotation, field) for name, field in RequirementAnalysis.model_fields.items()},
)


class BatchRequirementAnalysis(BaseModel):
    results: List[RequirementAnalysisItem] = Field(
        description="One analysis entry per generated requirement given, in the same order, with requirement_id matching the input id."
    )


# ============================================================
# Inter-rater agreement configuration (human sheet column -> LLM field name)
# ============================================================
HUMAN_TO_FIELD = {
    "Syntactic adherence": "syntactic_adherence",
    "Semantic adherence": "semantic_adherence",
    "Key concept manually matched to": "key_concept_manually_matched_to",
    "Secondary concept match": "secondary_concept_match",
    "Details match": "details_match",
    "Abstraction level": "abstraction_level",
    "Detail Level": "detail_level",
}

# Fields compared directly, one value per requirement
SIMPLE_AGREEMENT_FIELDS = [
    "syntactic_adherence",
    "semantic_adherence",
    "abstraction_level",
    "detail_level",
]

# Fields only meaningful per matched gold-standard id
PER_ID_AGREEMENT_FIELDS = ["secondary_concept_match", "details_match"]

PairList = List[Tuple[str, str]]


# ============================================================
# Analysis instructions - verbatim excerpt from the manual coding protocol.
# Sections "3 Key concept code class", "9 Recommended addition", "1 Purpose
# and preparation" and "4 Recording and review" are omitted: they either
# require source-code access not given to the LLM here, or are purely about
# human record-keeping/versioning and don't apply to a single LLM judgment.
# Table-formatted decision rules have been rendered as plain "Label: rule"
# lines; wording is otherwise unchanged from the protocol.
# ============================================================
ANALYSIS_INSTRUCTIONS = """
Coding procedure 
Unit of analysis. One generated requirement is one record, identified by project, generation variant and requirement ID. Secondary-concept and detail-match ratings apply separately to each generated–gold requirement pair. Many-to-many mappings between gold standard and generated requirements are permitted. 
Read each generated requirement in full and identify its primary functionality or user goal. Assess the nine variables in the codebook order, which follows the manual analysis workbook. Examine the full project gold standard requirement set when mapping requirements and the relevant implementation when identifying code classes. 
Assess secondary concepts and details only for mapped pairs. Assess abstraction and detail level for both matched and unmatched requirements. Evaluate recommended additions for unmatched functionality. Preserve the generated text and record the decisions separately.  
Concept decomposition. Before assigning the comparison variables, separate the generated requirement into: (1) its key concept, defined by the main domain action and its primary object or intended outcome; (2) secondary concepts, consisting of specified actors/roles and additional domain actions; and (3) details, consisting of data, values, formats, constraints, conditions, validations and prescribed steps. Do not use differences assigned to one category to determine another. 
For a mapped requirement, use the mapped gold requirement(s) as the primary reference when assessing functional scope. For an unmatched requirement, use the typical functional scope of the complete project gold set. 
When a decision remains ambiguous after applying the rules, do not infer missing information. Record the competing interpretations and retain the item for review rather than resolving the ambiguity through an unstated assumption. 

Coding criteria 
The variable names and sequence below follow Sheet1 of the manual analysis workbook. Use the allowed labels consistently and keep the assessments of presentation, meaning, scope and detail separate. 

1 Syntactic adherence 
Definition. Adherence to the presentation and syntactic format of the project’s gold-standard requirements. 
Output and applicability. Yes or No for every generated requirement. 
Decision rules. Record Yes when the same general structural pattern is clearly present, including relevant headings, sentence structure and style. Equivalent heading names are acceptable. Record No when the pattern is absent or substantially different. General readability alone does not establish adherence. 
Example. If the gold standard uses named use cases with actors and a flow, an equivalent structured use case receives Yes; a plain prose sentence receives No. For a gold standard written as prose, prose is not itself grounds for No. 

2 Semantic adherence 
Definition. Whether the generated requirement reflects the same domain-level semantic focus as the project’s gold-standard requirements. 
Output and applicability. Yes or No for every generated requirement. 
Decision rules. Record Yes when the requirement describes a domain-relevant goal, behavior, capability, rule, or constraint of the kind represented in the gold standard. Record No when it instead describes an implementation-specific or technical operation outside that scope. Technical details may be included without causing No if the main intent remains domain-relevant. 
Example. “The employee views book information” receives Yes because it describes domain-relevant behavior. “The system loads the book object into a session” receives No because it describes an internal implementation operation rather than domain-level behavior. 

4 Key concept manually matched to 
Definition. The gold-standard requirement(s) that share the generated requirement’s core functionality or user goal. Examples of key concepts include login, user deletion, and password recovery. 
Output and applicability. Unique gold-standard IDs, separated by commas, for every generated requirement with a match. Record no match when none applies. 
Decision rules. Examine the complete gold set and identify the core functionality expressed by the generated requirement. Match every gold requirement that expresses the same functionality, even when the requirements differ in scope, abstraction, actor, or detail. These differences are assessed separately. 
A broader generated requirement may therefore match several narrower gold requirements. Likewise, a requirement describing a substantive stage of a broader gold requirement may be matched when that stage directly contributes to the same functional goal. Do not match requirements based only on shared terminology, data, prerequisites, or incidental supporting steps. 
Example. “Search users” does not match “Delete users” merely because searching may be a step in deletion. However, a requirement describing the deletion of a selected user can match “Delete users” even if it covers only one stage or variant of the full gold-standard requirement. 

5 Secondary concept match 
Definition. Coverage of the mapped gold requirement’s secondary domain concepts: its specified actors or roles and additional user or domain actions beyond the shared key concept. 
Output and applicability. Yes, Partly, or No for each mapped gold ID, in the same order as the ID list. Not applicable for an unmatched requirement. 
Rating 
Yes: All specified secondary concepts are covered. 
Partly: At least one secondary concept is covered and at least one is missing or different. 
No: None of the specified secondary concepts are covered. 
Decision rules. Treat the mapped gold requirement as fixed; do not reconsider the key-concept match or abstraction level in this step. Identify its explicitly specified actors/roles and additional domain actions beyond the shared key concept, and compare these as discrete items with the generated requirement. Equivalent wording counts, but a generic or implied actor is not automatically equivalent to a specified role. Differences in abstraction or scope do not by themselves determine the rating; they are assessed separately under Abstraction level. Technical mechanisms, formats, validations, and other implementation details are assessed under Details match. 
Example. Gold standard requirement: a customer books a delivery and selects a preferred delivery time. If the shared key concept is booking a delivery, a generated requirement covering both the customer and the time selection receives Yes. One covering the customer and booking but omitting the time selection receives Partly. One describing the booking but with a different specified actor and no time selection receives No. 

6 Details match 
Definition. Preservation of the mapped gold requirement’s specified data fields, values, formats, constraints, conditional behavior and prescribed steps. 
Output and applicability. Yes, Partly or No for each mapped gold ID, in the same order as the ID list. Not applicable for an unmatched requirement. 
Rating 
Yes: All specified details are explicitly preserved. 
Partly: Nearly all details are preserved, with minor omissions or differences. Defining constraints and required behavior remain intact. 
No: Defining details are omitted, generalised or changed, or most specific information is absent. 
Decision rules. Treat the mapped gold standard requirement and its key and secondary concepts as fixed. Assess only the specific information used to qualify or constrain that functionality. Do not infer unstated details. A generic statement such as “validate input” does not preserve a specified validation rule. Additional generated details do not compensate for missing gold details. Differences in functionality, actors, or abstraction are assessed under the corresponding variables rather than under Details match. 
Examples. Gold standard requirement: “The customer enters a delivery date in YYYY-MM-DD format, and the date must be at least two days in the future.” A generated requirement preserving both the format and date constraint receives Yes. Preserving the date constraint but omitting the format receives Partly. Saying only that “the delivery date is validated” receives No. 

7 Abstraction level 
Definition. The breadth of the main functionality relative to the typical functional scope of the project’s gold-standard requirements. 
Output and applicability. Equal, Lower or Higher for every generated requirement, whether matched or unmatched. 
Rating 
Equal: Covers a complete goal or group of related actions of comparable scope. 
Lower: Covers a narrower action, supporting operation or internal subtask. 
Higher: Combines several goals or functions into a broader requirement. 
Decision rules. For a mapped requirement, compare the breadth of its main functionality with the mapped gold requirement(s). Record Lower when it expresses only a stage, subfunction, exception, or supporting operation of a gold-level goal; Equal when it represents approximately the same functional unit; and Higher when it combines multiple goals that are represented separately in the gold standard. For an unmatched requirement, make the same comparison against the typical functional unit in the complete gold set. Count domain functionality rather than fields, validations, technical steps, or amount of text. 
Example. If a typical gold requirement groups viewing, editing and deleting user data, deletion alone is Lower. Descriptions of internal system operations would be Lower. If the project usually specifies those actions separately, deletion may be Equal. Combining several separately specified goals into broader user management may be Higher. 

8 Detail Level 
Definition. The amount and specificity of concrete information relative to typical requirements in the complete project gold set. 
Output and applicability. Equal, Lower or Higher for every generated requirement, whether matched or unmatched. 
Rating 
Equal: Provides a comparable amount and specificity of concrete information. 
Lower: Provides fewer or less specific details about data, steps, conditions or outcomes. 
Higher: Provides more extensive concrete information. 
Decision rules. Assess the whole requirement against the project baseline. Concrete implementation details count; text length, repeated information, generic steps and technical vocabulary alone do not establish higher detail. Keep the amount of detail separate from whether those details match gold. 
Example. Where gold typically specifies the fields used to create a user, “create a user” has Lower detail. A requirement naming comparable fields may be Equal. A requirement adding extensive concrete validation and outcome rules may be Higher. 
A requirement may have Higher detail and Details match = No.  
"""


# ============================================================
# Load gold standard requirements for a dataset
# ============================================================
def load_gold_standard(excel_path: Path, dataset: str) -> pd.DataFrame:
    df = pd.read_excel(excel_path, sheet_name=f"{dataset}_gs", header=0)
    df.columns = [str(c).strip() for c in df.columns]
    return df[["#id", "Requirement"]].dropna(subset=["#id"]).reset_index(drop=True)


def format_gold_standard(df_gs: pd.DataFrame) -> str:
    return "\n\n".join(f"{row['#id']}: {row['Requirement']}" for _, row in df_gs.iterrows())


# ============================================================
# Load the generated requirements for one dataset/variant sheet
# ============================================================
def load_variant_requirements(excel_path: Path, dataset: str, variant: str) -> pd.DataFrame:
    sheet = f"{dataset}_{variant}"
    df_raw = pd.read_excel(excel_path, sheet_name=sheet, header=None)
    header = df_raw.iloc[1].astype(str).str.strip()
    df = df_raw.iloc[2:].copy()
    df.columns = header
    df = df.loc[:, ~(df.isna().all(axis=0))]
    df.columns = [str(c).strip() for c in df.columns]
    id_col = df.columns[0]
    req_col = "Requirement" if "Requirement" in df.columns else df.columns[1]
    df = df.dropna(subset=[id_col]).reset_index(drop=True)
    return df[[id_col, req_col]].rename(columns={id_col: "id", req_col: "requirement"})


# ============================================================
# Inter-rater agreement: load the human-authored sheet for one dataset/variant
# ============================================================
def load_human_analysis(excel_path: Path, dataset: str, variant: str) -> pd.DataFrame:
    sheet = f"{dataset}_{variant}"
    df_raw = pd.read_excel(excel_path, sheet_name=sheet, header=None)
    header = df_raw.iloc[1].astype(str).str.strip()
    df = df_raw.iloc[2:].copy()
    df.columns = header
    df = df.loc[:, ~(df.isna().all(axis=0))]
    df.columns = [str(c).strip() for c in df.columns]
    id_col = df.columns[0]
    df = df.rename(columns={id_col: "id", **HUMAN_TO_FIELD})
    df = df.dropna(subset=["id"]).reset_index(drop=True)
    df["id"] = df["id"].astype(str).str.strip()
    return df


def _norm(value) -> Optional[str]:
    """Normalize a cell value for comparison; returns None for empty/NA."""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    s = str(value).strip().lower()
    if s in ("", "nan", "na", "n/a", "none", "-"):
        return None
    return s


def split_list(value) -> List[str]:
    if _norm(value) is None:
        return []
    return [tok.strip() for tok in str(value).split(",") if tok.strip()]


def normalize_id(value) -> str:
    return str(value).strip().upper()


def jaccard_similarity(a: frozenset, b: frozenset) -> float:
    """Intersection-over-union of two gold-id sets. Both empty is treated as
    perfect trivial agreement (1.0), since both raters agree nothing maps."""
    if not a and not b:
        return 1.0
    return len(a & b) / len(a | b)


def dice_f1(a: frozenset, b: frozenset) -> float:
    """Set-level F1 (Dice coefficient): 2|A n B| / (|A| + |B|)."""
    if not a and not b:
        return 1.0
    if not a or not b:
        return 0.0
    return 2 * len(a & b) / (len(a) + len(b))


SetPairList = List[Tuple[frozenset, frozenset]]


def collect_agreement_pairs(
    dataset: str,
    variant: str,
    records: Dict[str, "PairList"],
    set_pairs: SetPairList,
    gold_ids: "frozenset[str]",
    run: int = 1,
) -> int:
    """Compare the human sheet against the LLM checkpoint for one dataset/variant/run,
    appending (human, llm) value pairs to `records`, keyed by field name, and
    (human_set, llm_set) gold-id set pairs to `set_pairs`.
    Returns the number of requirements analyzed by both raters."""
    try:
        human_df = load_human_analysis(EXCEL_PATH, dataset, variant)
    except Exception as e:
        print(f"Skipping {dataset}_{variant} (human sheet): {e}")
        return 0

    llm_checkpoint = load_checkpoint_for_run(dataset, variant, run)
    if not llm_checkpoint:
        print(f"Skipping {dataset}_{variant} run {run}: no LLM checkpoint found.")
        return 0
    llm_checkpoint_by_id = {normalize_id(req_id): row for req_id, row in llm_checkpoint.items()}

    n_pairs = 0
    for _, h_row in human_df.iterrows():
        l_row = llm_checkpoint_by_id.get(normalize_id(h_row["id"]))
        if l_row is None:
            continue
        n_pairs += 1

        for field in SIMPLE_AGREEMENT_FIELDS:
            h_val = _norm(h_row.get(field))
            l_val = _norm(l_row.get(field))
            if h_val is None or l_val is None:
                continue
            records.setdefault(field, []).append((h_val, l_val))

        h_ids = split_list(h_row.get("key_concept_manually_matched_to"))
        l_ids = split_list(l_row.get("key_concept_manually_matched_to"))
        # Gold-standard IDs are identifiers, not categorical labels. Normalize
        # their spelling before comparing them so capitalization/outer whitespace
        # cannot create a false disagreement.
        h_ids_norm = [normalize_id(i) for i in h_ids]
        l_ids_norm = [normalize_id(i) for i in l_ids]
        records.setdefault("is_mapped_to_gold", []).append(
            ("yes" if h_ids else "no", "yes" if l_ids else "no")
        )
        set_pairs.append((frozenset(h_ids_norm), frozenset(l_ids_norm)))

        # Strict per-gold-id binary agreement: for every gold-standard id, did each
        # rater include it in this requirement's matched-to set? Flattening this
        # into yes/no pairs (one per gold id) lets Cohen's kappa be computed over
        # the full binary indicator matrix, exactly-matching-ids-as-true style.
        h_ids_norm_set = set(h_ids_norm)
        l_ids_norm_set = set(l_ids_norm)
        for gold_id in gold_ids:
            records.setdefault("key_concept_match_strict", []).append((
                "yes" if gold_id in h_ids_norm_set else "no",
                "yes" if gold_id in l_ids_norm_set else "no",
            ))

        # Per-gold-id fields (secondary concept / details match) are only meaningful
        # for gold ids both raters mapped this requirement to.
        for field in PER_ID_AGREEMENT_FIELDS:
            h_labels = split_list(h_row.get(field))
            l_labels = split_list(l_row.get(field))
            if len(h_labels) != len(h_ids_norm) or len(l_labels) != len(l_ids_norm):
                print(
                    f"Skipping {field} for {dataset}_{variant} requirement {h_row['id']} "
                    f"run {run}: the number of labels does not match the number of mapped IDs."
                )
                continue
            h_map = dict(zip(h_ids_norm, h_labels))
            l_map = dict(zip(l_ids_norm, l_labels))
            for gold_id in set(h_map) & set(l_map):
                h_val, l_val = _norm(h_map[gold_id]), _norm(l_map[gold_id])
                if h_val is None or l_val is None:
                    continue
                records.setdefault(field, []).append((h_val, l_val))

    return n_pairs


def kappa_interpretation(k: float) -> str:
    # Landis & Koch (1977) rule of thumb.
    if k < 0:
        return "poor"
    if k < 0.20:
        return "slight"
    if k < 0.40:
        return "fair"
    if k < 0.60:
        return "moderate"
    if k < 0.80:
        return "substantial"
    return "almost perfect"


def summarize_agreement(records: Dict[str, "PairList"]) -> pd.DataFrame:
    rows = []
    for field, pairs in records.items():
        if not pairs:
            continue
        human_vals, llm_vals = zip(*pairs)
        n = len(pairs)
        percent_agreement = sum(h == l for h, l in pairs) / n * 100
        labels = sorted(set(human_vals) | set(llm_vals))
        kappa = cohen_kappa_score(human_vals, llm_vals, labels=labels) if len(labels) >= 2 else float("nan")
        rows.append({
            "field": field,
            "n_pairs": n,
            # Keep full precision here. Rounding belongs in the presentation
            # layer; otherwise means and SDs across runs are calculated from
            # already-rounded values.
            "percent_agreement": percent_agreement,
            "cohens_kappa": kappa if kappa == kappa else None,
            "interpretation": kappa_interpretation(kappa) if kappa == kappa else "n/a",
        })
    return pd.DataFrame(rows)


def summarize_set_agreement(set_pairs: SetPairList) -> pd.DataFrame:
    """Cohen's kappa doesn't apply to multi-label sets like
    key_concept_manually_matched_to (there's no fixed small label space to build
    a confusion matrix from). Instead, report the standard set-agreement metrics:
    mean Jaccard index (IoU) and mean Dice/set-F1, per requirement, macro-averaged.
    Both-empty (neither rater mapped the requirement to any gold id) counts as
    trivial perfect agreement (1.0)."""
    n = len(set_pairs)
    if n == 0:
        return pd.DataFrame()
    exact_match_pct = sum(h == l for h, l in set_pairs) / n * 100
    mean_jaccard = sum(jaccard_similarity(h, l) for h, l in set_pairs) / n
    mean_dice_f1 = sum(dice_f1(h, l) for h, l in set_pairs) / n
    return pd.DataFrame([{
        "field": "key_concept_manually_matched_to",
        "n_pairs": n,
        # Retain full precision until results are displayed or written.
        "exact_set_match_pct": exact_match_pct,
        "mean_jaccard": mean_jaccard,
        "mean_dice_f1": mean_dice_f1,
    }])


# Human-readable labels for the internal field names, used only for pretty-printing
# console output. Fields not listed here fall back to their raw name.
FIELD_DISPLAY_NAMES = {
    "syntactic_adherence": "Adheres syntactically",
    "semantic_adherence": "Adheres semantically",
    "abstraction_level": "Level of abstraction",
    "detail_level": "Level of detail",
    "is_mapped_to_gold": "Is matched to gold standard",
    "key_concept_match_strict": "Key concept match (strict, per gold id)",
    "secondary_concept_match": "Secondary concepts match",
    "details_match": "Details match",
    "key_concept_manually_matched_to": "Key concept match (set overlap)",
}


def format_agreement_table(df: Optional[pd.DataFrame]) -> Optional[pd.DataFrame]:
    """Builds a compact, human-readable view of a summarize_agreement()-style dataframe
    for console printing: pretty-prints the variable name (via FIELD_DISPLAY_NAMES,
    falling back to the raw field name) and keeps only Variable/Agreement (%)/Cohen's
    kappa, plus any run/variant/n_runs grouping columns present."""
    if df is None or df.empty:
        return df
    out = df.copy()
    out["field"] = out["field"].map(lambda f: FIELD_DISPLAY_NAMES.get(f, f))
    lead_cols = [c for c in ("run", "variant", "n_runs") if c in out.columns]
    if "percent_agreement_sd" in out.columns:
        out["agreement_mean_sd"] = out.apply(
            lambda row: _format_mean_sd(row["percent_agreement"], row["percent_agreement_sd"], 1), axis=1
        )
        out["kappa_mean_sd"] = out.apply(
            lambda row: _format_mean_sd(row["cohens_kappa"], row["cohens_kappa_sd"], 3), axis=1
        )
        out = out[lead_cols + ["field", "agreement_mean_sd", "kappa_mean_sd"]]
        return out.rename(columns={
            "field": "Variable",
            "agreement_mean_sd": "Agreement (%), mean ± SD",
            "kappa_mean_sd": "Cohen's κ, mean ± SD",
        })
    out["percent_agreement"] = out["percent_agreement"].round(1)
    out["cohens_kappa"] = out["cohens_kappa"].round(3)
    out = out[lead_cols + ["field", "percent_agreement", "cohens_kappa"]]
    return out.rename(columns={
        "field": "Variable",
        "percent_agreement": "Agreement (%)",
        "cohens_kappa": "Cohen's \u03ba",
    })


def format_set_agreement_table(df: Optional[pd.DataFrame]) -> Optional[pd.DataFrame]:
    """Same idea as format_agreement_table(), for summarize_set_agreement()-style
    dataframes: keeps Variable/exact set match (%)/mean Jaccard only."""
    if df is None or df.empty:
        return df
    out = df.copy()
    out["field"] = out["field"].map(lambda f: FIELD_DISPLAY_NAMES.get(f, f))
    lead_cols = [c for c in ("run", "variant", "n_runs") if c in out.columns]
    if "exact_set_match_pct_sd" in out.columns:
        out["exact_mean_sd"] = out.apply(
            lambda row: _format_mean_sd(row["exact_set_match_pct"], row["exact_set_match_pct_sd"], 1), axis=1
        )
        out["jaccard_mean_sd"] = out.apply(
            lambda row: _format_mean_sd(row["mean_jaccard"], row["mean_jaccard_sd"], 3), axis=1
        )
        out = out[lead_cols + ["field", "exact_mean_sd", "jaccard_mean_sd"]]
        return out.rename(columns={
            "field": "Variable",
            "exact_mean_sd": "Exact set match (%), mean ± SD",
            "jaccard_mean_sd": "Jaccard, mean ± SD",
        })
    out["exact_set_match_pct"] = out["exact_set_match_pct"].round(1)
    out["mean_jaccard"] = out["mean_jaccard"].round(3)
    out = out[lead_cols + ["field", "exact_set_match_pct", "mean_jaccard"]]
    return out.rename(columns={
        "field": "Variable",
        "exact_set_match_pct": "Exact set match (%)",
        "mean_jaccard": "Mean Jaccard",
    })


def _format_mean_sd(mean: float, sd: float, decimals: int) -> str:
    """Format a mean and run-to-run sample SD for the average worksheet."""
    if pd.isna(mean):
        return "n/a"
    if pd.isna(sd):
        return f"{mean:.{decimals}f} ± n/a"
    return f"{mean:.{decimals}f} ± {sd:.{decimals}f}"


def average_agreement_summaries(
    run_summaries: List[pd.DataFrame], group_cols: Optional[List[str]] = None
) -> pd.DataFrame:
    """Averages a list of per-run summarize_agreement() dataframes (one per run) into
    a single dataframe with mean and sample SD (ddof=1) for percent agreement
    and Cohen's kappa per field (or per group_cols, e.g. ["variant", "field"]),
    plus n_runs = how many runs contributed. Row order follows first appearance."""
    group_cols = group_cols or ["field"]
    non_empty = [df for df in run_summaries if df is not None and not df.empty]
    if not non_empty:
        return pd.DataFrame()
    combined = pd.concat(non_empty, ignore_index=True)
    order = combined[group_cols].drop_duplicates()
    agg = combined.groupby(group_cols, sort=False).agg(
        n_runs=("percent_agreement", "size"),
        n_pairs=("n_pairs", "mean"),
        percent_agreement=("percent_agreement", "mean"),
        percent_agreement_sd=("percent_agreement", "std"),
        cohens_kappa=("cohens_kappa", "mean"),
        cohens_kappa_sd=("cohens_kappa", "std"),
    ).reset_index()
    agg["n_pairs"] = agg["n_pairs"].round(1)
    agg["percent_agreement"] = agg["percent_agreement"].round(1)
    agg["percent_agreement_sd"] = agg["percent_agreement_sd"].round(1)
    agg["cohens_kappa"] = agg["cohens_kappa"].round(3)
    agg["cohens_kappa_sd"] = agg["cohens_kappa_sd"].round(3)
    agg["interpretation"] = agg["cohens_kappa"].apply(lambda k: kappa_interpretation(k) if pd.notna(k) else "n/a")
    return order.merge(agg, on=group_cols, how="left")


def average_set_agreement_summaries(
    run_set_summaries: List[pd.DataFrame], group_cols: Optional[List[str]] = None
) -> pd.DataFrame:
    """Same idea as average_agreement_summaries, for summarize_set_agreement() outputs."""
    group_cols = group_cols or ["field"]
    non_empty = [df for df in run_set_summaries if df is not None and not df.empty]
    if not non_empty:
        return pd.DataFrame()
    combined = pd.concat(non_empty, ignore_index=True)
    order = combined[group_cols].drop_duplicates()
    agg = combined.groupby(group_cols, sort=False).agg(
        n_runs=("n_pairs", "size"),
        n_pairs=("n_pairs", "mean"),
        exact_set_match_pct=("exact_set_match_pct", "mean"),
        exact_set_match_pct_sd=("exact_set_match_pct", "std"),
        mean_jaccard=("mean_jaccard", "mean"),
        mean_jaccard_sd=("mean_jaccard", "std"),
        mean_dice_f1=("mean_dice_f1", "mean"),
        mean_dice_f1_sd=("mean_dice_f1", "std"),
    ).reset_index()
    agg["n_pairs"] = agg["n_pairs"].round(1)
    agg["exact_set_match_pct"] = agg["exact_set_match_pct"].round(1)
    agg["exact_set_match_pct_sd"] = agg["exact_set_match_pct_sd"].round(1)
    agg["mean_jaccard"] = agg["mean_jaccard"].round(3)
    agg["mean_jaccard_sd"] = agg["mean_jaccard_sd"].round(3)
    agg["mean_dice_f1"] = agg["mean_dice_f1"].round(3)
    agg["mean_dice_f1_sd"] = agg["mean_dice_f1_sd"].round(3)
    return order.merge(agg, on=group_cols, how="left")


AGREEMENT_SHEET_PREFIX = "Agreement"


def _write_titled_table(ws, title: str, df: pd.DataFrame, start_row: int) -> int:
    """Writes `title` as a bold header at `start_row` (1-indexed openpyxl row) and
    the dataframe starting the row after. Returns the next free row (1-indexed)."""
    cell = ws.cell(row=start_row, column=1, value=title)
    cell.font = Font(bold=True)
    row = start_row + 1
    for values in dataframe_to_rows(df, index=False, header=True):
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        row += 1
    return row + 1  # blank gap row


def _write_agreement_sheet(wb, sheet_name: str, index: int, title_suffix: str, results: Dict[str, Optional[pd.DataFrame]]) -> None:
    ws = wb.create_sheet(sheet_name, index)
    row = 1
    row = _write_titled_table(ws, f"Inter-rater agreement (human vs LLM){title_suffix}", format_agreement_table(results["summary"]), row)
    if results.get("set_summary") is not None and not results["set_summary"].empty:
        row = _write_titled_table(ws, f"Key concept match agreement (set overlap){title_suffix}", format_set_agreement_table(results["set_summary"]), row)
    if results.get("variant_summary") is not None:
        row = _write_titled_table(ws, f"Per-variant inter-rater agreement{title_suffix}", format_agreement_table(results["variant_summary"]), row)
    if results.get("variant_set_summary") is not None:
        row = _write_titled_table(ws, f"Per-variant key concept match agreement (set overlap){title_suffix}", format_set_agreement_table(results["variant_set_summary"]), row)


def write_agreement_to_workbook(
    per_run_results: Dict[int, Dict[str, Optional[pd.DataFrame]]],
    avg_results: Dict[str, Optional[pd.DataFrame]],
) -> None:
    """Writes one 'Agreement_run{N}' sheet per run plus a final 'Agreement_avg' sheet
    with the results averaged across runs, all placed first in
    Manual_analysis_llm_protocol_openai.xlsx. Written directly with openpyxl (rather
    than repeated pandas to_excel calls) so that writing one table doesn't wipe out
    the tables already written to the same sheet."""
    wb = load_workbook(OUTPUT_XLSX) if OUTPUT_XLSX.exists() else Workbook()
    for name in list(wb.sheetnames):
        if name.startswith(AGREEMENT_SHEET_PREFIX):
            del wb[name]
    if wb.sheetnames == ["Sheet"] and wb["Sheet"].max_row == 1 and wb["Sheet"].max_column == 1 and wb["Sheet"]["A1"].value is None:
        # Freshly created workbook still has the default empty sheet; drop it.
        del wb["Sheet"]

    index = 0
    for run in sorted(per_run_results.keys()):
        sheet_name = _excel_sheet_name(AGREEMENT_SHEET_PREFIX + "_run", str(run))
        _write_agreement_sheet(wb, sheet_name, index, f" - Run {run}", per_run_results[run])
        index += 1
    _write_agreement_sheet(wb, f"{AGREEMENT_SHEET_PREFIX}_avg", index, " - Mean ± sample SD across runs", avg_results)

    wb.save(OUTPUT_XLSX)
    print(f"\nWrote agreement results to '{AGREEMENT_SHEET_PREFIX}_run*' and '{AGREEMENT_SHEET_PREFIX}_avg' sheets in {OUTPUT_XLSX}")


def compute_agreement(datasets: List[str], variants: List[str], per_variant: bool = False) -> None:
    # Compare human vs LLM separately for every run found in the checkpoints (see
    # --runs), then average kappa/percent-agreement across runs. This reflects
    # whatever runs actually exist on disk, regardless of how many --runs were
    # requested in this particular invocation.
    runs = sorted({r for dataset in datasets for variant in variants for r in discover_runs(dataset, variant)})
    if not runs:
        print("No LLM checkpoints found; nothing to compute agreement for.")
        return
    print(f"Found checkpoint data for run(s): {runs}")

    per_run_results: Dict[int, Dict[str, Optional[pd.DataFrame]]] = {}
    per_run_summaries: List[pd.DataFrame] = []
    per_run_set_summaries: List[pd.DataFrame] = []
    per_run_variant_summaries: List[pd.DataFrame] = []
    per_run_variant_set_summaries: List[pd.DataFrame] = []

    for run in runs:
        overall_records: Dict[str, "PairList"] = {}
        per_variant_records: Dict[str, Dict[str, "PairList"]] = {}
        overall_set_pairs: SetPairList = []
        per_variant_set_pairs: Dict[str, SetPairList] = {}
        total_pairs = 0

        for dataset in datasets:
            try:
                gold_ids = frozenset(normalize_id(i) for i in load_gold_standard(EXCEL_PATH, dataset)["#id"])
            except Exception as e:
                print(f"Skipping {dataset}: failed to load gold standard ids: {e}")
                continue
            for variant in variants:
                local_records: Dict[str, "PairList"] = {}
                local_set_pairs: SetPairList = []
                n = collect_agreement_pairs(dataset, variant, local_records, local_set_pairs, gold_ids, run=run)
                total_pairs += n
                for field, pairs in local_records.items():
                    overall_records.setdefault(field, []).extend(pairs)
                    per_variant_records.setdefault(variant, {}).setdefault(field, []).extend(pairs)
                overall_set_pairs.extend(local_set_pairs)
                per_variant_set_pairs.setdefault(variant, []).extend(local_set_pairs)

        print(f"\n--- Run {run}: matched {total_pairs} requirements with both a human and LLM analysis. ---")

        summary = summarize_agreement(overall_records)
        print(format_agreement_table(summary).to_string(index=False))
        per_run_summaries.append(summary)

        set_summary = summarize_set_agreement(overall_set_pairs)
        if not set_summary.empty:
            print(format_set_agreement_table(set_summary).to_string(index=False))
            per_run_set_summaries.append(set_summary)

        variant_summary = None
        variant_set_summary = None
        if per_variant:
            rows = []
            for variant, records in per_variant_records.items():
                df = summarize_agreement(records)
                df.insert(0, "variant", variant)
                rows.append(df)
            if rows:
                variant_summary = pd.concat(rows, ignore_index=True)
                per_run_variant_summaries.append(variant_summary)

            set_rows = []
            for variant, pairs in per_variant_set_pairs.items():
                df = summarize_set_agreement(pairs)
                if df.empty:
                    continue
                df.insert(0, "variant", variant)
                set_rows.append(df)
            if set_rows:
                variant_set_summary = pd.concat(set_rows, ignore_index=True)
                per_run_variant_set_summaries.append(variant_set_summary)

        per_run_results[run] = {
            "summary": summary,
            "set_summary": set_summary,
            "variant_summary": variant_summary,
            "variant_set_summary": variant_set_summary,
        }

    avg_summary = average_agreement_summaries(per_run_summaries)
    print(f"\n=== Mean ± sample SD of inter-rater agreement across {len(runs)} run(s) (human vs LLM) ===")
    print(format_agreement_table(avg_summary).to_string(index=False))

    avg_set_summary = average_set_agreement_summaries(per_run_set_summaries)
    if not avg_set_summary.empty:
        print(f"\n=== Mean ± sample SD of key concept match agreement across {len(runs)} run(s) (set overlap) ===")
        print(format_set_agreement_table(avg_set_summary).to_string(index=False))

    avg_variant_summary = None
    avg_variant_set_summary = None
    if per_variant:
        if per_run_variant_summaries:
            avg_variant_summary = average_agreement_summaries(per_run_variant_summaries, group_cols=["variant", "field"])
            print("\n=== Per-variant inter-rater agreement (mean ± sample SD across runs) ===")
            print(format_agreement_table(avg_variant_summary).to_string(index=False))
        if per_run_variant_set_summaries:
            avg_variant_set_summary = average_set_agreement_summaries(per_run_variant_set_summaries, group_cols=["variant", "field"])
            print("\n=== Per-variant key concept match agreement (set overlap, mean ± sample SD across runs) ===")
            print(format_set_agreement_table(avg_variant_set_summary).to_string(index=False))

    write_agreement_to_workbook(per_run_results, {
        "summary": avg_summary,
        "set_summary": avg_set_summary,
        "variant_summary": avg_variant_summary,
        "variant_set_summary": avg_variant_set_summary,
    })


# ============================================================
# Prompting
# ============================================================
def build_system_prompt(instructions: str, gold_standard: str) -> str:
    return f"""You are performing an analysis of an automatically generated software 
requirement, comparing it against a gold standard set of requirements for the same system, 
following the manual coding protocol excerpt below exactly.

{instructions}

Calibration note:
Treat the labels as categorical decisions, not confidence levels. Do not use Partly or Equal 
as fallback choices when uncertain. Assign a label only when its stated conditions are met. 
Do not infer actors, actions, details, or functionality that are not stated or clearly entailed 
by the requirement. Keep the assessment dimensions separate: structure, semantic focus, 
key functionality, secondary concepts, matching details, functional breadth, and amount of 
detail must be judged independently. When evidence is insufficient for a more inclusive 
interpretation, do not assume typical or default behavior.

Here are the GOLD STANDARD requirements for this system:
{gold_standard}

Analyze the generated requirement you are given and return your assessment for 
ALL of the properties above, following the protocol's decision rules exactly."""


def build_user_prompt(requirement_id: str, requirement_text: str) -> str:
    return f"Generated requirement {requirement_id}: {requirement_text}"


def analyze_requirement(
    client: OpenAI,
    model: str,
    system_prompt: str,
    requirement_id: str,
    requirement_text: str,
) -> Optional[RequirementAnalysis]:
    user_prompt = build_user_prompt(requirement_id, requirement_text)
    for attempt in range(1, RETRIES + 1):
        try:
            response = client.chat.completions.parse(
                model=model,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
                response_format=RequirementAnalysis,
            )
            return response.choices[0].message.parsed
        except Exception as e:
            print(f"  LLM call failed for {requirement_id} (attempt {attempt}/{RETRIES}): {e}")
            if attempt < RETRIES:
                time.sleep(RETRY_BACKOFF_SECONDS)
    return None


def build_batch_user_prompt(requirements: List[Tuple[str, str]]) -> str:
    parts = [f"Generated requirement {req_id}: {req_text}" for req_id, req_text in requirements]
    return (
        "You are given the following generated requirements to analyze. Return exactly "
        "one result per requirement, in the same order, with requirement_id set to the "
        "id given below (e.g. 'code_req_42').\n\n" + "\n\n".join(parts)
    )


def analyze_requirements_batch(
    client: OpenAI,
    model: str,
    system_prompt: str,
    requirements: List[Tuple[str, str]],
) -> Dict[str, RequirementAnalysisItem]:
    """Analyzes a whole batch (e.g. all pending requirements for a dataset/variant) in
    a single LLM call. Returns a dict keyed by requirement_id."""
    user_prompt = build_batch_user_prompt(requirements)
    for attempt in range(1, RETRIES + 1):
        try:
            response = client.chat.completions.parse(
                model=model,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
                response_format=BatchRequirementAnalysis,
            )
            parsed = response.choices[0].message.parsed
            return {item.requirement_id: item for item in parsed.results}
        except Exception as e:
            print(f"  Batch LLM call failed (attempt {attempt}/{RETRIES}): {e}")
            if attempt < RETRIES:
                time.sleep(RETRY_BACKOFF_SECONDS)
    return {}


# ============================================================
# Checkpointing (resume support)
# ============================================================
def serialize_result(result: RequirementAnalysis) -> Dict[str, str]:
    """Flatten the (possibly list-valued) analysis fields into comma-joined strings,
    matching the format used in the human-authored sheets (e.g. 'EA1, EA2')."""
    data = result.model_dump()
    flat = {}
    for field in ANALYSIS_FIELDS:
        value = data[field]
        if isinstance(value, list):
            flat[field] = ", ".join(str(v) for v in value)
        else:
            flat[field] = value
    return flat


def checkpoint_path(dataset: str, variant: str) -> Path:
    return CHECKPOINT_DIR / f"{dataset}_{variant}.csv"


# Checkpoint CSV columns. "run" tags which repeat (1, 2, 3, ...) a row came from,
# so multiple runs can be appended to the same file/output workbook and later
# compared/averaged (see --runs and compute_agreement).
CHECKPOINT_FIELDNAMES = ["run", "id", "requirement"] + ANALYSIS_FIELDS


def _read_checkpoint_df(dataset: str, variant: str) -> Optional[pd.DataFrame]:
    path = checkpoint_path(dataset, variant)
    if not path.exists():
        return None
    df = pd.read_csv(path, dtype=str)
    if "run" not in df.columns:
        # Checkpoint predates the "run" column; treat all existing rows as run 1.
        df.insert(0, "run", "1")
    return df


def discover_runs(dataset: str, variant: str) -> List[int]:
    """Returns the sorted list of distinct run numbers present in this
    dataset/variant's checkpoint file (empty if no checkpoint exists)."""
    df = _read_checkpoint_df(dataset, variant)
    if df is None or df.empty:
        return []
    return sorted({int(r) for r in df["run"]})


def load_checkpoint_for_run(dataset: str, variant: str, run: int) -> Dict[str, dict]:
    df = _read_checkpoint_df(dataset, variant)
    if df is None:
        return {}
    df = df[df["run"].astype(int) == run]
    return {row["id"]: row.to_dict() for _, row in df.iterrows()}


def _migrate_checkpoint_run_column(path: Path) -> None:
    """If an existing checkpoint file predates the "run" column, rewrite it with
    run=1 for all existing rows so new rows can be appended with a consistent header."""
    if not path.exists():
        return
    with open(path, "r", newline="", encoding="utf-8") as f:
        header = f.readline().strip().split(",")
    if header and header[0] == "run":
        return
    df = pd.read_csv(path, dtype=str)
    df.insert(0, "run", "1")
    df.to_csv(path, index=False)


def append_checkpoint_row(dataset: str, variant: str, run: int, row: dict) -> None:
    path = checkpoint_path(dataset, variant)
    _migrate_checkpoint_run_column(path)
    is_new = not path.exists()
    with open(path, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CHECKPOINT_FIELDNAMES)
        if is_new:
            writer.writeheader()
        writer.writerow({"run": run, **row})


# ============================================================
# Processing one dataset/variant sheet
# ============================================================
def process_variant_batch(
    client: OpenAI,
    model: str,
    dataset: str,
    variant: str,
    df_reqs: pd.DataFrame,
    system_prompt: str,
    run: int = 1,
    limit: Optional[int] = None,
    batch_size: Optional[int] = None,
) -> None:
    """Like process_variant, but sends multiple requirements to the LLM per call
    instead of one at a time. By default all pending requirements for this
    dataset/variant/run are sent in a single call; pass batch_size to chunk them.
    `run` identifies this repeat measurement (see --runs); requirements already
    completed for this specific run are skipped, but the same requirement is
    re-analyzed independently for each other run."""
    already_done = load_checkpoint_for_run(dataset, variant, run)
    pending = [
        (str(req.id).strip(), str(req.requirement))
        for req in df_reqs.itertuples(index=False)
        if str(req.id).strip() not in already_done
    ]
    if limit is not None:
        pending = pending[:limit]
    if not pending:
        print(f"[{dataset}/{variant}] run {run}: nothing to do (batch mode).")
        return

    chunks = [pending[i:i + batch_size] for i in range(0, len(pending), batch_size)] if batch_size else [pending]
    for chunk in chunks:
        print(f"[{dataset}/{variant}] run {run}: analyzing batch of {len(chunk)} requirements...")
        results_by_id = analyze_requirements_batch(client, model, system_prompt, chunk)
        if not results_by_id:
            print(f"  Batch call failed after {RETRIES} attempts, skipping this batch.")
            continue
        for req_id, req_text in chunk:
            result = results_by_id.get(req_id)
            if result is None:
                print(f"  Missing result for {req_id} in batch response, skipping.")
                continue
            row = {"id": req_id, "requirement": req_text, **serialize_result(result)}
            append_checkpoint_row(dataset, variant, run, row)


# ============================================================
# Compile all checkpoints into the final output workbook
# ============================================================
def _excel_sheet_name(base: str, suffix: str) -> str:
    """Builds an Excel-safe (<=31 char) sheet name as "{base}_{suffix}", truncating
    `base` (not `suffix`) if needed so the suffix (e.g. a run number) stays intact."""
    suffix_str = f"_{suffix}"
    max_base_len = max(31 - len(suffix_str), 1)
    return base[:max_base_len] + suffix_str


def compile_output_workbook() -> None:
    """Writes one sheet per dataset/variant/run (e.g. "eANCI_Feature_RAG_GPT_1",
    "..._2", ...) so results from different runs can be inspected side by side
    instead of being merged into a single sheet."""
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        wrote_any = False
        for dataset in DATASETS:
            for variant in VARIANTS:
                for run in discover_runs(dataset, variant):
                    checkpoint = load_checkpoint_for_run(dataset, variant, run)
                    if not checkpoint:
                        continue
                    df = pd.DataFrame(list(checkpoint.values()))
                    df = df[["id", "requirement"] + [f for f in ANALYSIS_FIELDS if f in df.columns]]
                    sheet_name = _excel_sheet_name(f"{dataset}_{variant}", str(run))
                    # Match the human sheet layout: title row, then header, then data.
                    title_row = pd.DataFrame([[f"llm_generated_{dataset}_{variant}_run{run}"] + [None] * (len(df.columns) - 1)],
                                              columns=df.columns)
                    out_df = pd.concat([title_row, df], ignore_index=True)
                    out_df.to_excel(writer, sheet_name=sheet_name, header=True, index=False)
                    wrote_any = True
        if not wrote_any:
            # openpyxl/pandas require at least one sheet to be written.
            pd.DataFrame().to_excel(writer, sheet_name="Sheet", index=False)
    print(f"Wrote combined results to {OUTPUT_XLSX}")


# ============================================================
# Main
# ============================================================
def main():
    parser = argparse.ArgumentParser(description="Run automated (LLM-based) manual analysis using the verbatim coding protocol as instructions.")
    parser.add_argument("--datasets", nargs="*", default=DATASETS, help="Subset of datasets to process.")
    parser.add_argument("--variants", nargs="*", default=VARIANTS, help="Subset of variants to process.")
    parser.add_argument("--runs", type=int, default=1, help="Number of times to repeat the whole LLM analysis per dataset/variant (e.g. --runs 3). All runs are appended to the same checkpoint/output file, tagged with a 'run' column, and --agreement reports per-run results plus their mean and sample SD.")
    parser.add_argument("--limit", type=int, default=None, help="Max number of new requirements per variant (for testing).")
    parser.add_argument("--compile-only", action="store_true", help="Skip LLM calls, just rebuild the output workbook from checkpoints.")
    parser.add_argument("--batch-size", type=int, default=None, help="Max number of requirements per LLM call (default: all pending requirements in one call).")
    parser.add_argument("--agreement", action="store_true", help="After running (or with --compile-only, using existing checkpoints), also compute Cohen's kappa between the human and LLM analyses, per run and averaged across runs.")
    parser.add_argument("--agreement-only", action="store_true", help="Skip the LLM pipeline entirely and only compute agreement from existing checkpoints.")
    parser.add_argument("--per-variant", action="store_true", help="With --agreement, also break down agreement by variant.")
    args = parser.parse_args()

    if args.agreement_only:
        compile_output_workbook()
        compute_agreement(args.datasets, args.variants, per_variant=args.per_variant)
        return

    CHECKPOINT_DIR.mkdir(parents=True, exist_ok=True)

    if not args.compile_only:
        client = get_client()
        model = get_model_name()

        for dataset in args.datasets:
            print(f"=== Loading context for {dataset} ===")
            df_gs = load_gold_standard(EXCEL_PATH, dataset)
            gold_standard_text = format_gold_standard(df_gs)
            system_prompt = build_system_prompt(ANALYSIS_INSTRUCTIONS, gold_standard_text)

            for variant in args.variants:
                try:
                    df_reqs = load_variant_requirements(EXCEL_PATH, dataset, variant)
                except Exception as e:
                    print(f"Failed to load sheet {dataset}_{variant}: {e}")
                    continue
                for run in range(1, args.runs + 1):
                    process_variant_batch(client, model, dataset, variant, df_reqs, system_prompt, run=run, limit=args.limit, batch_size=args.batch_size)


    compile_output_workbook()

    if args.agreement:
        compute_agreement(args.datasets, args.variants, per_variant=args.per_variant)


if __name__ == "__main__":
    main()
